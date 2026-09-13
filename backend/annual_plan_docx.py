"""DOCX renderer for the structured full-year Scheme of Work."""
from __future__ import annotations
from io import BytesIO
from typing import Any
from annual_plan_documents import StructuredAnnualPlanDocument
from lesson_plan_docx import EMERALD, TEAL, SLATE, LIGHT_SLATE, _add_bullets, _add_footer_branding, _keep_row_together, _repeat_table_header, _section_heading, _set_cell_margins, _set_cell_shading, _set_cell_text

def render_structured_annual_plan_docx(document: StructuredAnnualPlanDocument, *, teacher: str | None = None, meta: dict[str, Any] | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Mm, Pt, RGBColor
    meta = meta or {}
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Mm(15); section.bottom_margin = Mm(15); section.left_margin = Mm(15); section.right_margin = Mm(15)
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(9.5)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.add_run("  |  ".join(item for item in (teacher, str(meta.get("schoolName") or "").strip()) if item))
    _add_footer_branding(footer, meta, logo_height=Inches(.24), right_edge=section.page_width-section.left_margin-section.right_margin)
    title = doc.add_paragraph(); title_run = title.add_run(document.title); title_run.bold = True; title_run.font.name = "Arial"; title_run.font.size = Pt(21); title_run.font.color.rgb = RGBColor.from_string(EMERALD)
    details = [("SUBJECT", document.subject), ("LEVEL / YEAR", document.level), ("CLASS / GROUP", document.class_context), ("ACADEMIC YEAR", document.academic_year)]
    details = [(label, value) for label, value in details if value]
    table = doc.add_table(rows=1, cols=len(details)); table.autofit = False
    for index, (label, value) in enumerate(details):
        cell = table.cell(0, index); _set_cell_shading(cell, LIGHT_SLATE); _set_cell_margins(cell)
        _set_cell_text(cell, f"{label}\n{value}", bold=False, size=8.5)
    _keep_row_together(table.rows[0])
    for heading, items in (("Planning basis", document.planning_basis), ("Calendar constraints", document.calendar_constraints)):
        if items: _section_heading(doc, heading); _add_bullets(doc, items)
    _section_heading(doc, "Term by term course map")
    table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"; table.autofit = False
    headings = ("Dates", "Topic / chapter", "Learning focus", "Practical / assessment evidence")
    widths = (1.0, 1.35, 2.7, 1.65)
    for index, (heading, width) in enumerate(zip(headings, widths)):
        cell = table.rows[0].cells[index]; cell.width = Inches(width); _set_cell_shading(cell, TEAL); _set_cell_text(cell, heading, bold=True, size=8, colour="FFFFFF")
    _repeat_table_header(table.rows[0])
    for item in document.course_map:
        row = table.add_row()
        for index, value in enumerate((item.dates, item.topic, item.learning_focus, item.practical_or_assessment or "-")):
            row.cells[index].width = Inches(widths[index]); _set_cell_text(row.cells[index], value, bold=index < 2, size=8.2)
        _keep_row_together(row)
    for heading, items in (("Teaching and assessment rhythm", document.teaching_assessment_rhythm), ("Practical or project programme", document.practical_project_programme), ("Progress checkpoints revision and recovery buffers", document.checkpoints_and_buffers), ("Planning sources", document.planning_sources), ("Assumptions and adjustment rules", document.assumptions_and_adjustment_rules)):
        if items: _section_heading(doc, heading); _add_bullets(doc, items)
    if document.first_lesson: _section_heading(doc, "Optional detailed first lesson"); doc.add_paragraph(document.first_lesson)
    result = BytesIO(); doc.save(result); return result.getvalue()
