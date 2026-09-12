from __future__ import annotations
import os, shutil, sys, tempfile, unittest
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from pathlib import Path
from uuid import uuid4
from unittest.mock import patch
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.orm.query import Query
from sqlalchemy.pool import StaticPool

BACKEND = Path(__file__).resolve().parents[1]; sys.path.insert(0, str(BACKEND))
os.environ.setdefault("JWT_SECRET", "aac-document-test-secret-0123456789")
import main, models
from db import Base


class AacDocumentEndpointTests(unittest.TestCase):
    def setUp(self):
        # Python 3.14's Windows TemporaryDirectory private ACL cannot be re-entered
        # by this sandboxed interpreter.  A unique normal directory still exercises
        # the real upload writes and is removed by tearDown.
        self.tmp = Path(tempfile.gettempdir()) / f"aac-endpoint-{uuid4().hex}"; self.tmp.mkdir()
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine); self.Session = sessionmaker(bind=self.engine); self.db = self.Session()
        self.owner = models.UserModel(email="pfitzgerald@preskilkenny.ie", password_hash="x", role="teacher", is_active=True, email_verified=True)
        self.other = models.UserModel(email="other@example.test", password_hash="x", role="teacher", is_active=True, email_verified=True)
        self.db.add_all([self.owner, self.other]); self.db.flush()
        self.cls = models.ClassModel(owner_user_id=self.owner.id, name="Fictional Physics", subject="Physics", aac_planner_enabled=True)
        self.other_cls = models.ClassModel(owner_user_id=self.other.id, name="Other", subject="Physics", aac_planner_enabled=True)
        self.db.add_all([self.cls, self.other_cls]); self.db.flush()
        self.project = models.AacProjectModel(class_id=self.cls.id, owner_user_id=self.owner.id, title="Fictional AAC", subject="Physics")
        self.other_project = models.AacProjectModel(class_id=self.other_cls.id, owner_user_id=self.other.id, title="Other AAC", subject="Physics")
        self.db.add_all([self.project, self.other_project]); self.db.commit()
        main.app.dependency_overrides[main.get_db] = lambda: self.db
        main.app.dependency_overrides[main.get_current_user] = lambda: self.owner
        self.client = TestClient(main.app)
        self.upload_patch = patch.object(main, "UPLOADS_DIR", self.tmp); self.upload_patch.start()

    def tearDown(self):
        self.upload_patch.stop(); main.app.dependency_overrides.clear(); self.db.close(); Base.metadata.drop_all(self.engine); self.engine.dispose(); shutil.rmtree(self.tmp)

    @staticmethod
    def docx_bytes():
        from docx import Document
        doc = Document(); doc.add_paragraph("Fictional deadline is 20 April 2027."); stream = BytesIO(); doc.save(stream); return stream.getvalue()

    def upload(self, **extra):
        if not self.db.query(models.AacPlanRevisionModel).filter_by(project_id=self.project.id, state="draft").first():
            self.add_approvable_draft()
        data = {"purpose": "specification", **extra}
        return self.client.post(f"/classes/{self.cls.id}/aac/documents", data=data, files={"file": ("fictional.docx", self.docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})

    def source(self, purpose, filename, sections, state="extracted"):
        row = models.AacSourceDocumentModel(project_id=self.project.id, owner_user_id=self.owner.id, purpose=purpose, display_filename=filename, storage_key=f"aac/test/{uuid4().hex}_{filename}", content_type="application/pdf", size_bytes=12, sha256="a" * 64, extraction_state=state, extracted_sections_json=sections)
        self.db.add(row); self.db.flush()
        draft = self.db.query(models.AacPlanRevisionModel).filter_by(project_id=self.project.id, state="draft").first() or self.add_approvable_draft()
        draft.source_document_ids_json = [*(draft.source_document_ids_json or []), row.id]
        self.db.commit(); return row

    def add_approvable_draft(self, version=1, stage_id="fictional-stage", stage_name="Fictional stage"):
        inputs = {"weekly_minutes":30,"fifth_year_end":"2026-05-29","sixth_year_restart":"2026-09-14","controlling_deadline":"2027-04-20","internal_completion_target":"2027-04-06","official_deadline_confirmed":True,"sixth_year_calendar_status":"reviewed"}
        revision = models.AacPlanRevisionModel(project_id=self.project.id, version=version, state="draft", source_requirements_json=[], assumptions_json=[], source_document_ids_json=[], planning_inputs_json=inputs, plan_json={"stages":[{"id":stage_id,"name":stage_name,"estimated_minutes":30,"completion_date":"2027-03-20"}]})
        self.db.add(revision); self.db.commit(); return revision

    def review_token(self):
        return self.client.get(f"/classes/{self.cls.id}/aac").json()["project"]["revision"]["review_token"]

    def test_owner_uploads_lists_private_document_without_storage_key(self):
        self.add_approvable_draft()
        created = self.upload(); self.assertEqual(created.status_code, 200, created.text); self.assertNotIn("storage_key", created.json())
        listed = self.client.get(f"/classes/{self.cls.id}/aac/documents"); self.assertEqual(listed.status_code, 200); self.assertEqual(len(listed.json()), 1); self.assertIn("paragraph 1", listed.json()[0]["sections"][0]["reference"])

    def test_fresh_draft_has_no_historical_active_sources_then_persists_upload(self):
        prior = self.add_approvable_draft()
        historical = self.source("specification", "historical.pdf", [{"reference": "page 1", "text": "Historical source"}])
        prior.source_document_ids_json = [historical.id]; self.db.commit()
        response = self.client.post(f"/classes/{self.cls.id}/aac/new-draft", json={"title": "Fresh", "subject": "Art", "weekly_minutes": 30, "current_year_stage": "sixth_year"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/documents").json(), [])
        uploaded = self.upload()
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        listed = self.client.get(f"/classes/{self.cls.id}/aac/documents")
        self.assertEqual(listed.status_code, 200)
        self.assertEqual([item["id"] for item in listed.json()], [uploaded.json()["id"]])
        self.db.expire_all()
        self.assertEqual(self.db.query(models.AacPlanRevisionModel).filter_by(id=prior.id).one().source_document_ids_json, [historical.id])

    def test_new_tracker_keeps_historical_metadata_until_clean_setup_is_saved(self):
        prior = self.add_approvable_draft()
        old_title, old_subject = self.project.title, self.project.subject
        started = self.client.post(f"/classes/{self.cls.id}/aac/new-draft")
        self.assertEqual(started.status_code, 200, started.text)
        clean = started.json()["revision"]
        self.assertTrue(clean["plan"]["tracker_setup_pending"])
        self.assertEqual(clean["source_document_ids"], [])
        self.db.expire_all()
        retained = self.db.query(models.AacPlanRevisionModel).filter_by(id=prior.id).one()
        self.assertEqual(retained.plan_json["tracker_metadata"]["title"], old_title)
        self.assertEqual(retained.plan_json["tracker_metadata"]["subject"], old_subject)
        configured = self.client.put(f"/classes/{self.cls.id}/aac/tracker-details", json={"title": "New tracker", "subject": "Art", "weekly_minutes": 30, "current_year_stage": "sixth_year"})
        self.assertEqual(configured.status_code, 200, configured.text)
        self.assertFalse(configured.json()["revision"]["plan"]["tracker_setup_pending"])
        self.assertEqual(configured.json()["title"], "New tracker")

    def test_only_named_reviewers_can_discover_or_use_draft_pilot(self):
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac").status_code, 200)
        main.app.dependency_overrides[main.get_current_user] = lambda: self.other
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac").status_code, 403)
        self.assertEqual(self.client.put(f"/classes/{self.cls.id}/aac/enabled", json={"enabled": True}).status_code, 403)
        self.assertEqual(self.client.post(f"/classes/{self.cls.id}/aac/projects", json={"title": "No access", "subject": "Physics", "weekly_minutes": 30, "current_year_stage": "fifth_year"}).status_code, 403)

    def test_pilot_rejects_provider_proposals_and_calendar_publication(self):
        document = self.upload().json(); revision = self.db.query(models.AacPlanRevisionModel).filter_by(project_id=self.project.id, state="draft").one()
        with patch.object(main, "_enforce_ai_feature_limit") as enforce:
            proposal = self.client.post(f"/classes/{self.cls.id}/aac/proposals", json={"document_ids": [document["id"]], "planning_inputs": {}})
        self.assertEqual(proposal.status_code, 409); enforce.assert_not_called()
        approval = self.client.post(f"/classes/{self.cls.id}/aac/approve", json={"revision_id": revision.id, "review_token": self.review_token()})
        self.assertEqual(approval.status_code, 409); self.assertEqual(self.db.query(models.CalendarEvent).count(), 0)

    def test_deadline_review_is_read_only_and_uses_only_specifications(self):
        draft = models.AacPlanRevisionModel(project_id=self.project.id, version=1, state="draft", source_requirements_json=[], assumptions_json=[], source_document_ids_json=[], planning_inputs_json={"controlling_deadline": ""}, plan_json={"stages": [{"id": "teacher-stage", "name": "Teacher text"}]})
        self.db.add(draft); self.db.commit()
        specification = self.source("specification", "fictional-brief.pdf", [{"reference": "page 2", "text": "The coursework must be carried out and submitted to the class teacher by Friday, 11 December 2026."}])
        self.source("fifth_year_calendar", "calendar.pdf", [{"reference": "page 1", "text": "School administration submits forms by 20 January 2030."}])
        response = self.client.get(f"/classes/{self.cls.id}/aac/deadline-candidates")
        self.assertEqual(response.status_code, 200, response.text); body = response.json()
        self.assertEqual(body["status"], "candidate"); self.assertEqual(len(body["candidates"]), 1)
        self.assertEqual(body["candidates"][0]["date"], "2026-12-11"); self.assertEqual(body["candidates"][0]["source_document_id"], specification.id)
        self.db.expire_all(); unchanged = self.db.query(models.AacPlanRevisionModel).filter_by(id=draft.id).one()
        self.assertEqual(unchanged.plan_json["stages"][0]["name"], "Teacher text"); self.assertEqual(unchanged.planning_inputs_json["controlling_deadline"], "")

    def test_deadline_review_reports_calendar_only_missing_and_conflicting_states(self):
        self.source("sixth_year_calendar", "calendar.pdf", [{"reference": "page 1", "text": "Classes resume 1 September 2027."}])
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/deadline-candidates").json()["status"], "specification_needed")
        self.source("specification", "no-deadline.pdf", [{"reference": "page 1", "text": "Published 1 January 2027."}])
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/deadline-candidates").json()["status"], "no_clear_deadline")
        self.source("specification", "conflict.pdf", [{"reference": "page 2", "text": "Students must hand in coursework by 11 December 2026."}, {"reference": "page 3", "text": "Students must hand in coursework by 12 December 2026."}])
        conflict = self.client.get(f"/classes/{self.cls.id}/aac/deadline-candidates").json()
        self.assertEqual(conflict["status"], "conflicting_candidates"); self.assertEqual({item["date"] for item in conflict["candidates"]}, {"2026-12-11", "2026-12-12"})

    def test_owner_removes_unapproved_source_and_private_file(self):
        document = self.upload().json(); row = self.db.query(models.AacSourceDocumentModel).filter_by(id=document["id"]).one(); path = self.tmp / row.storage_key
        response = self.client.delete(f"/classes/{self.cls.id}/aac/documents/{row.id}")
        self.assertEqual(response.status_code, 200, response.text); self.assertTrue(response.json()["removed"]); self.assertFalse(path.exists()); self.assertIsNone(self.db.query(models.AacSourceDocumentModel).filter_by(id=row.id).first())

    def test_owner_removes_safe_missing_file_record_and_invalidates_derived_evidence(self):
        document = self.upload().json(); row = self.db.query(models.AacSourceDocumentModel).filter_by(id=document["id"]).one(); path = self.tmp / row.storage_key
        revision = models.AacPlanRevisionModel(project_id=self.project.id, version=1, state="draft", source_document_ids_json=[row.id], source_requirements_json=[{"source_document_id": row.id, "source_reference": "paragraph 1", "text": "Derived evidence"}], assumptions_json=[], plan_json={"stages": [{"id": "teacher-stage", "name": "Keep teacher stage"}], "candidate_deadlines": [{"source_document_id": row.id, "source_reference": "paragraph 1", "date": "2027-04-20"}]}, planning_inputs_json={"official_deadline_confirmed": True, "official_deadline_source": {"source_document_id": row.id}})
        self.db.query(models.AacPlanRevisionModel).filter_by(project_id=self.project.id, state="draft").delete()
        self.db.add(revision); self.db.commit(); path.unlink()
        response = self.client.delete(f"/classes/{self.cls.id}/aac/documents/{row.id}")
        self.assertEqual(response.status_code, 200, response.text); self.assertTrue(response.json()["removed"]); self.assertTrue(response.json()["stored_file_already_absent"])
        self.assertIn("already absent", response.json()["message"]); self.assertIsNone(self.db.query(models.AacSourceDocumentModel).filter_by(id=row.id).first())
        self.db.expire_all(); saved = self.db.query(models.AacPlanRevisionModel).filter_by(id=revision.id).one()
        self.assertEqual(saved.plan_json["stages"][0]["name"], "Keep teacher stage"); self.assertEqual(saved.plan_json["candidate_deadlines"], []); self.assertEqual(saved.source_requirements_json, []); self.assertEqual(saved.source_document_ids_json, []); self.assertFalse(saved.planning_inputs_json["official_deadline_confirmed"])

    def test_other_teacher_and_unsafe_path_cannot_remove_source(self):
        document = self.upload().json(); row = self.db.query(models.AacSourceDocumentModel).filter_by(id=document["id"]).one(); path = self.tmp / row.storage_key
        main.app.dependency_overrides[main.get_current_user] = lambda: self.other
        self.assertIn(self.client.delete(f"/classes/{self.cls.id}/aac/documents/{row.id}").status_code, (403,404)); self.assertTrue(path.exists())
        main.app.dependency_overrides[main.get_current_user] = lambda: self.owner; row.storage_key = "../outside"; self.db.commit()
        self.assertEqual(self.client.delete(f"/classes/{self.cls.id}/aac/documents/{row.id}").status_code, 409); self.assertIsNotNone(self.db.query(models.AacSourceDocumentModel).filter_by(id=row.id).first())

    def test_remove_commit_failure_restores_private_file_and_record(self):
        document = self.upload().json(); row = self.db.query(models.AacSourceDocumentModel).filter_by(id=document["id"]).one(); path = self.tmp / row.storage_key
        with patch.object(self.db, "commit", side_effect=RuntimeError("fictional commit failure")):
            response = self.client.delete(f"/classes/{self.cls.id}/aac/documents/{row.id}")
        self.assertEqual(response.status_code, 503); self.assertTrue(path.exists()); self.db.rollback(); self.assertIsNotNone(self.db.query(models.AacSourceDocumentModel).filter_by(id=row.id).first()); self.assertEqual(list(path.parent.glob(path.name + ".deleting-*")), [])

    def test_student_or_other_teacher_cannot_read_or_upload(self):
        main.app.dependency_overrides[main.get_current_user] = lambda: self.other
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/documents").status_code, 403)
        self.assertEqual(self.upload().status_code, 403)
        main.app.dependency_overrides.pop(main.get_current_user)
        # QR/PIN sessions do not satisfy the teacher identity dependency.
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/documents").status_code, 401)

    def test_bad_documents_and_commit_failure_leave_no_orphan_or_replacement(self):
        good = self.upload(); self.assertEqual(good.status_code, 200); old = self.db.query(models.AacSourceDocumentModel).one(); old_path = self.tmp / old.storage_key
        response = self.client.post(f"/classes/{self.cls.id}/aac/documents", data={"purpose":"specification"}, files={"file": ("bad.txt", b"bad", "text/plain")})
        self.assertEqual(response.status_code, 422); self.assertTrue(old_path.exists()); self.assertEqual(old.extraction_state, "extracted")
        with patch.object(self.db, "commit", side_effect=RuntimeError("fictional db failure")):
            response = self.upload()
        self.assertEqual(response.status_code, 503); self.db.rollback(); self.assertEqual(self.db.query(models.AacSourceDocumentModel).count(), 1); self.assertTrue(old_path.exists()); self.assertEqual(len(list((self.tmp / "aac" / str(self.project.id)).glob("*"))), 1)

    def test_cleanup_failure_preserves_the_original_safe_persistence_error(self):
        self.add_approvable_draft()
        with patch.object(self.db, "commit", side_effect=RuntimeError("fictional db failure")), patch.object(main, "UPLOADS_DIR", self.tmp / "cleanup-failure"):
            with patch.object(main, "_remove_aac_source_file", side_effect=OSError("fictional cleanup failure")):
                response = self.upload()
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["detail"], "The document could not be saved. Please try again.")
        self.db.rollback(); self.assertEqual(self.db.query(models.AacSourceDocumentModel).count(), 0)

    def test_draft_save_does_not_publish_milestones(self):
        payload = {"source_requirements":[],"plan":{"stages":[{"id":"draft-only","name":"Draft only","estimated_minutes":30,"completion_date":"2027-03-20"}]},"assumptions":[],"source_document_ids":[],"planning_inputs":{"weekly_minutes":30}}
        response = self.client.put(f"/classes/{self.cls.id}/aac/revision", json=payload)
        self.assertEqual(response.status_code, 200, response.text); self.assertEqual(self.db.query(models.CalendarEvent).count(), 0)

    def test_both_reviewers_are_limited_to_their_own_classes(self):
        self.other.email = "dcampion@preskilkenny.ie"; self.db.commit()
        for user, own, foreign in ((self.owner, self.cls, self.other_cls), (self.other, self.other_cls, self.cls)):
            main.app.dependency_overrides[main.get_current_user] = lambda: user
            self.assertEqual(self.client.get(f"/classes/{own.id}/aac").status_code, 200)
            for method, suffix, payload in (("get", "", None), ("post", "/new-draft", None), ("put", "/tracker-details", {"title":"Fictional", "subject":"Art"}), ("get", "/documents", None)):
                response = self.client.request(method, f"/classes/{foreign.id}/aac{suffix}", json=payload)
                self.assertEqual(response.status_code, 404, response.text)

    def test_complete_source_review_save_reload_recalculation_and_fresh_history(self):
        from docx import Document
        doc = Document(); doc.add_paragraph("Students must complete and hand in their AAC by 12 March 2027.")
        doc.add_paragraph("Stage 1: Research\nStage 2: Investigation\nStage 3: Report")
        stream = BytesIO(); doc.save(stream)
        with patch.object(self, "docx_bytes", return_value=stream.getvalue()):
            uploaded = self.upload()
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        source_id = uploaded.json()["id"]
        review = self.client.get(f"/classes/{self.cls.id}/aac/deadline-candidates").json()
        self.assertEqual(review["candidates"][0]["date"], "2027-03-12")
        self.assertEqual(review["candidates"][0]["source_reference"], "paragraph 1")
        self.assertEqual(len(review["stage_structure"]), 3)
        inputs = {"weekly_minutes":30,"planned_start":"2026-09-16","normal_finish_target":"2027-02-08","final_classroom_deadline":"2027-02-26","controlling_deadline":"2027-03-12","fifth_year_end":"2026-05-29","sixth_year_restart":"2026-09-14","official_deadline_confirmed":True}
        stages = [{"id":f"source-{source_id}-{item['number']}","name":item["name"],"estimated_minutes":45,"completion_date":"2026-11-02"} for item in review["stage_structure"]]
        payload = {"plan":{"stages":stages},"planning_inputs":inputs,"source_document_ids":[source_id]}
        saved = self.client.put(f"/classes/{self.cls.id}/aac/revision", json=payload)
        self.assertEqual(saved.status_code, 200, saved.text)
        revision_id = saved.json()["revision"]["id"]
        reloaded = self.client.get(f"/classes/{self.cls.id}/aac").json()["project"]["revision"]
        self.assertEqual(reloaded["plan"]["stages"], stages)
        proposed = self.client.get(f"/classes/{self.cls.id}/aac?recalculate_dates=true").json()["project"]["revision"]
        self.assertTrue(all(s["proposed_completion_date"] >= inputs["planned_start"] for s in proposed["schedule"]["stages"]))
        self.db.expire_all()
        self.assertEqual(self.db.get(models.AacPlanRevisionModel, revision_id).plan_json["stages"], stages)
        self.assertEqual(self.client.post(f"/classes/{self.cls.id}/aac/new-draft").status_code, 200)
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/documents").json(), [])
        self.assertEqual(self.client.get(f"/classes/{self.cls.id}/aac/deadline-candidates").json()["status"], "specification_needed")
        self.assertTrue(self.client.get(f"/classes/{self.cls.id}/aac").json()["project"]["revision"]["plan"]["tracker_setup_pending"])
        replacement = self.upload(); self.assertEqual(replacement.status_code, 200, replacement.text)
        self.db.expire_all()
        old_source = self.db.get(models.AacSourceDocumentModel, source_id)
        self.assertEqual(old_source.extraction_state, "extracted")
        self.assertEqual(self.client.delete(f"/classes/{self.cls.id}/aac/documents/{source_id}").status_code, 200)
        self.assertTrue((self.tmp / old_source.storage_key).is_file())
        self.assertEqual(self.db.get(models.AacPlanRevisionModel, revision_id).source_document_ids_json, [source_id])
        self.assertEqual(self.db.query(models.CalendarEvent).count(), 0)
        self.assertEqual(self.db.query(models.AacStudentProgressModel).count(), 0)


if __name__ == "__main__": unittest.main()
