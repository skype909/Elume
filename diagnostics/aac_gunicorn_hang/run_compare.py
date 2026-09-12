"""Bounded Linux/Gunicorn comparison harness; never imports application code itself."""
from __future__ import annotations

import argparse
import concurrent.futures
import http.client
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
import shutil
import signal
import socket
import subprocess
import sys
import time
from uuid import uuid4
from progress_checks import check_progress, check_recovery

ROOT = Path(__file__).resolve().parents[2]
HARNESS = Path(__file__).resolve().parent
SECRET = "diagnostic-only-jwt-secret-not-used-outside-ci-7d6e5c4b3a2f"

def write(log: Path, text: str) -> None:
    log.parent.mkdir(parents=True, exist_ok=True)
    with log.open("a", encoding="utf-8") as stream: stream.write(datetime.now(timezone.utc).isoformat() + " " + text + "\n")

def request(port: int, method: str, path: str, *, token: str | None = None, body: bytes | None = None, content_type: str | None = None, timeout: float = 5) -> tuple[int, str]:
    conn = http.client.HTTPConnection("127.0.0.1", port, timeout=timeout); headers = {"Accept": "application/json"}
    if token: headers["Authorization"] = f"Bearer {token}"
    if content_type: headers["Content-Type"] = content_type
    try:
        conn.request(method, path, body=body, headers=headers); response = conn.getresponse()
        payload = response.read().decode("utf-8", "replace")
        if path == "/auth/me" and not token:
            if response.status != 401 or response.getheader("Content-Type", "").split(";")[0] != "application/json" or not isinstance(json.loads(payload), dict):
                raise AssertionError("auth probe did not return 401 application/json")
        return response.status, payload
    finally: conn.close()

def json_request(port: int, method: str, path: str, token: str, value: dict) -> tuple[int, dict]:
    status, payload = request(port, method, path, token=token, body=json.dumps(value).encode(), content_type="application/json")
    return status, json.loads(payload or "{}")

def multipart_upload(port: int, class_id: int, token: str) -> int:
    from docx import Document
    from io import BytesIO
    doc = Document(); doc.add_paragraph("Fictional coursework must be submitted to the class teacher by 12th March 2027."); doc.add_paragraph("Stage 1: Getting Started\nStage 2: Planning\nStage 3: Investigation")
    stream = BytesIO(); doc.save(stream); boundary = "----aac" + uuid4().hex
    body = b"".join((f"--{boundary}\r\nContent-Disposition: form-data; name=\"purpose\"\r\n\r\nspecification\r\n".encode(), f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"fictional-brief.docx\"\r\nContent-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n".encode(), stream.getvalue(), b"\r\n", f"--{boundary}--\r\n".encode()))
    return request(port, "POST", f"/classes/{class_id}/aac/documents", token=token, body=body, content_type=f"multipart/form-data; boundary={boundary}")[0]

def seed(source: Path, database_url: str) -> tuple[int, str]:
    if database_url != "postgresql+psycopg2://aac_diagnostic:diagnostic-only-not-a-secret@127.0.0.1:5432/aac_diagnostic":
        raise ValueError("Only the fixed disposable CI database is permitted")
    script = """import os
os.environ['DATABASE_URL']=os.environ['DIAGNOSTIC_DATABASE_URL']; os.environ['JWT_SECRET']=os.environ['DIAGNOSTIC_JWT_SECRET']
from db import Base,engine,SessionLocal
import models
from jose import jwt
Base.metadata.drop_all(bind=engine); Base.metadata.create_all(bind=engine); db=SessionLocal()
user=models.UserModel(email='pfitzgerald@preskilkenny.ie',password_hash='x',role='teacher',is_active=True,email_verified=True); db.add(user); db.flush()
klass=models.ClassModel(owner_user_id=user.id,name='Synthetic AAC',subject='Physics',aac_planner_enabled=True); db.add(klass); db.flush()
project=models.AacProjectModel(class_id=klass.id,owner_user_id=user.id,title='Synthetic tracker',subject='Physics'); db.add(project); db.flush()
revision=models.AacPlanRevisionModel(project_id=project.id,version=1,state='draft',source_requirements_json=[],source_document_ids_json=[],assumptions_json=[],planning_inputs_json={'weekly_minutes':30,'planned_start':'2026-09-14','normal_finish_target':'2027-02-08','final_classroom_deadline':'2027-02-20','controlling_deadline':'2027-03-12','fifth_year_end':'2026-05-29','sixth_year_restart':'2026-09-14'},plan_json={'stages':[{'id':'synthetic-stage','name':'Synthetic stage','completion_date':'2026-11-01','estimated_minutes':30,'checkpoints':[]}]}); db.add(revision)
db.add_all([models.StudentModel(class_id=klass.id,first_name=n,active=True) for n in ['Synthetic Alice','Synthetic Bob']]); db.flush()
dan=models.UserModel(email='dcampion@preskilkenny.ie',password_hash='x',role='teacher',is_active=True,email_verified=True); outsider=models.UserModel(email='outsider@example.test',password_hash='x',role='teacher',is_active=True,email_verified=True); db.add_all([dan,outsider]); db.flush()
other=models.ClassModel(owner_user_id=dan.id,name='Synthetic Dan class',subject='Physics',aac_planner_enabled=True); db.add(other); db.flush()
db.add(models.StudentModel(class_id=other.id,first_name='Synthetic Dan student',active=True))
other_project=models.AacProjectModel(class_id=other.id,owner_user_id=dan.id,title='Synthetic Dan tracker',subject='Physics'); db.add(other_project); db.flush()
db.add(models.AacPlanRevisionModel(project_id=other_project.id,version=1,state='draft',plan_json={'stages':[{'id':'dan-stage','name':'Dan stage','completion_date':'2026-11-01'}]})); db.commit()
print(klass.id); print(jwt.encode({'sub':str(user.id)},os.environ['DIAGNOSTIC_JWT_SECRET'],algorithm='HS256'))"""
    env = {**os.environ, "DIAGNOSTIC_DATABASE_URL": database_url, "DIAGNOSTIC_JWT_SECRET": SECRET, "PYTHONDONTWRITEBYTECODE": "1"}
    result = subprocess.run([sys.executable, "-c", script], cwd=source / "backend", env=env, capture_output=True, text=True, timeout=45, check=True)
    class_id, token = result.stdout.strip().splitlines()[-2:]; return int(class_id), token

def dump_stacks(process: subprocess.Popen[bytes], log: Path) -> None:
    try: os.killpg(process.pid, signal.SIGUSR1)
    except ProcessLookupError: pass
    time.sleep(1); write(log, "DIAGNOSTIC: SIGUSR1 sent to the disposable Gunicorn process group for test-only faulthandler capture")

def run_revision(label: str, revision: str, port: int, workdir: Path, database_url: str) -> None:
    source = workdir / label; subprocess.run(["git", "worktree", "add", "--detach", str(source), revision], cwd=ROOT, check=True, stdout=subprocess.DEVNULL)
    log = workdir / "logs" / f"{label}.log"
    try:
        try:
            class_id, token = seed(source, database_url)
        except Exception as exc:
            write(log, f"RESULT:setup-failure:{type(exc).__name__}")
            raise RuntimeError("Disposable database seeding failed; see setup classification") from None
        env = {**os.environ, "DATABASE_URL": database_url, "JWT_SECRET": SECRET, "ELUME_UPLOADS_DIR": str(workdir / label / "uploads"), "PYTHONDONTWRITEBYTECODE": "1", "PYTHONUTF8": "1", "PYTHONPATH": os.pathsep.join((str(HARNESS), str(source / "backend")))}
        command = [sys.executable, "-m", "gunicorn", "main:app", "--workers", "1", "--worker-class", "uvicorn.workers.UvicornWorker", "--bind", f"127.0.0.1:{port}"]
        write(log, f"REVISION:{revision}\nRUNTIME:{sys.version.split()[0]}\nCOMMAND:gunicorn workers=1 uvicorn.workers.UvicornWorker")
        with log.open("ab") as stream: process = subprocess.Popen(command, cwd=source / "backend", env=env, stdout=stream, stderr=subprocess.STDOUT, start_new_session=True)
        write(log, f"MASTER_PID:{process.pid}")
        def auth_probe(timeout=3):
            started = time.monotonic()
            status, _ = request(port, "GET", "/auth/me", timeout=timeout)
            write(log, f"AUTH:status={status} content_type=application/json seconds={time.monotonic()-started:.3f}")
            return status
        try:
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                try:
                    if auth_probe(timeout=2) == 401 and "Application startup complete." in log.read_text(): break
                except (OSError, TimeoutError, socket.timeout): time.sleep(0.2)
            else: raise TimeoutError("startup auth probe did not return 401")
            worker_ids = re.findall(r"\[(\d+)\].*Application startup complete\.", log.read_text())
            if not worker_ids: raise AssertionError("No worker startup-complete PID found")
            write(log, f"WORKER_READY_PIDS:{','.join(worker_ids)}")
            if [auth_probe() for _ in range(8)] != [401] * 8: raise AssertionError("serial unauthenticated probes failed")
            with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
                probes = [pool.submit(auth_probe) for _ in range(8)]
                if label == "suspect":
                    if json_request(port, "POST", f"/classes/{class_id}/aac/new-draft", token, {})[0] != 200: raise AssertionError("fresh tracker failed")
                    if json_request(port, "PUT", f"/classes/{class_id}/aac/tracker-details", token, {"title":"Synthetic tracker","subject":"Physics","weekly_minutes":30,"current_year_stage":"sixth_year"})[0] != 200: raise AssertionError("tracker setup failed")
                if multipart_upload(port, class_id, token) != 200: raise AssertionError("fictional upload failed")
                status, raw = request(port, "GET", f"/classes/{class_id}/aac/deadline-candidates", token=token, timeout=5)
                review = json.loads(raw)
                if status != 200 or review.get("status") != "candidate" or len(review.get("stage_structure", [])) != 3: raise AssertionError("deadline/stage review failed")
                if review["candidates"][0]["date"] != "2027-03-12": raise AssertionError("extracted deadline mismatch")
                docs = json.loads(request(port, "GET", f"/classes/{class_id}/aac/documents", token=token)[1])
                stages = [{"id":f"reviewed-{item['number']}","name":item["name"],"estimated_minutes":90,"completion_date":"2026-11-02","checkpoints":[]} for item in review["stage_structure"]]
                inputs = {"weekly_minutes":30,"planned_start":"2026-09-16","normal_finish_target":"2027-02-08","final_classroom_deadline":"2027-02-26","controlling_deadline":"2027-03-12","fifth_year_end":"2026-05-29","sixth_year_restart":"2026-09-14","official_deadline_confirmed":True}
                status, saved = json_request(port, "PUT", f"/classes/{class_id}/aac/revision", token, {"plan":{"stages":stages},"planning_inputs":inputs,"source_document_ids":[item["id"] for item in docs]})
                if status != 200: raise AssertionError("draft save failed")
                reloaded = json.loads(request(port, "GET", f"/classes/{class_id}/aac", token=token)[1])["project"]["revision"]
                if reloaded["plan"]["stages"] != stages: raise AssertionError("saved minutes/dates did not persist")
                status, raw = request(port, "GET", f"/classes/{class_id}/aac?recalculate_dates=true", token=token, timeout=5)
                if status != 200: raise AssertionError("recalculation failed")
                if label == "suspect":
                    scheduled = json.loads(raw)["project"]["revision"]["schedule"]["stages"]
                    if not all(item.get("proposed_completion_date") for item in scheduled) or len(scheduled) != 3: raise AssertionError("replacement suggestions missing")
                unchanged = json.loads(request(port, "GET", f"/classes/{class_id}/aac", token=token)[1])["project"]["revision"]
                if unchanged["plan"]["stages"] != stages: raise AssertionError("recalculation mutated saved dates")
                if [future.result() for future in probes] != [401] * 8: raise AssertionError("concurrent unauthenticated probes failed")
                write(log, "AAC:upload=200 review=source-backed-three-stages save=200 reload=90-minutes recalculation=read-only concurrent-auth=8x401-json")
                if label == "suspect":
                    identity, progress = check_progress(port,class_id,token,request,json_request,auth_probe,log,write)
                    if json_request(port, "POST", f"/classes/{class_id}/aac/new-draft", token, {})[0] != 200: raise AssertionError("second fresh start failed")
                    if json.loads(request(port, "GET", f"/classes/{class_id}/aac/documents", token=token)[1]) != []: raise AssertionError("fresh sources not isolated")
                    write(log, "AAC:fresh-start=200 active-sources=empty")
                    fresh = json.loads(request(port,"GET",f"/classes/{class_id}/aac/students",token=token)[1])
                    assert fresh["tracker_id"] != identity and all(s["check_in"]["version"] == 0 for s in fresh["students"])
                    historical = json.loads(request(port,"GET",f"/classes/{class_id}/aac/students?tracker_id={identity}",token=token)[1])
                    assert historical["read_only"] and historical["students"][0]["check_in"] == progress
                    write(log,"PROGRESS:fresh-tracker=isolated prior-history=retained")
            if re.findall(r"\[(\d+)\].*Application startup complete\.", log.read_text()) != worker_ids: raise AssertionError("worker restarted during diagnostic")
            write(log, "RESULT:pass")
        except Exception as exc:
            dump_stacks(process, log); write(log, f"RESULT:fail:{type(exc).__name__}"); raise
        finally:
            if process.poll() is None: os.killpg(process.pid, signal.SIGTERM)
            try: process.wait(timeout=10)
            except subprocess.TimeoutExpired: os.killpg(process.pid, signal.SIGKILL)
    finally: subprocess.run(["git", "worktree", "remove", "--force", str(source)], cwd=ROOT, check=False, stdout=subprocess.DEVNULL)

def main() -> int:
    parser = argparse.ArgumentParser(); parser.add_argument("--healthy", required=True); parser.add_argument("--suspect", required=True); args = parser.parse_args()
    workdir = Path(os.environ["DIAGNOSTIC_WORKDIR"]).resolve()
    runner_temp = Path(os.environ["RUNNER_TEMP"]).resolve()
    if workdir.parent != runner_temp or workdir.name != "aac-gunicorn-diagnostic": raise ValueError("Unsafe diagnostic directory")
    if workdir.exists(): raise ValueError("Diagnostic directory already exists; automatic repeats are forbidden")
    (workdir / "logs").mkdir(parents=True)
    os.environ["DIAGNOSTIC_JWT_SECRET"] = SECRET
    run_revision("healthy", args.healthy, 18182, workdir, os.environ["DIAGNOSTIC_DATABASE_URL"])
    run_revision("suspect", args.suspect, 18183, workdir, os.environ["DIAGNOSTIC_DATABASE_URL"])
    from jose import jwt
    check_recovery(ROOT,HARNESS,args.healthy,workdir,os.environ["DIAGNOSTIC_DATABASE_URL"],jwt.encode({"sub":"1"},SECRET,algorithm="HS256"),request,write,dump_stacks)
    return 0

if __name__ == "__main__": raise SystemExit(main())
